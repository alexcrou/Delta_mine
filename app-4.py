# ============================================================
#  ГОРНЫЙ ДЕЛЬТА v3.0 — единый файл, без внешних модулей
#  Все разделы встроены: штаб, greenfield, мастер, документы,
#  чек-лист, ассистент, устав/ТЗ, DWG, подписка, владелец.
#  Переключатель светлой/тёмной темы в сайдбаре.
#  requirements.txt: streamlit, pypdf, python-docx, openpyxl
# ============================================================
import streamlit as st
import json, os, datetime, re

OWNER_PASSWORD = "CHANGE_ME_2026"
PRICE_BASIC, PRICE_PRO = 4900, 14900
PROMO_CODES = {"PILOT2026": "pro", "DEMO": "basic"}
LOG_FILE = "users_log.jsonl"
DOCS_DIR = "knowledge"
os.makedirs(DOCS_DIR, exist_ok=True)

st.set_page_config(page_title="Горный Дельта — ИИ-платформа РП", page_icon="⛏", layout="wide")

# ---------- СОСТОЯНИЕ ----------
S = st.session_state
defaults = {
    "plan": "free", "show_owner": False, "wizard_done": False, "project": None,
    "docs": [], "theme": "light",
    "gf_project": {}, "gf_tasks": [], "gf_requirements": [],
    "hq_facts": {},
    "roadmap": [],   # [{work, start, finish, pct, criterion}]
    "budget": [],    # [{item, y1, y2, total}]
    "risks": [],     # [{risk, prob, owner}]
}
for k, v in defaults.items():
    if k not in S:
        S[k] = v

# ---------- ТЕМЫ (светлая / тёмная) ----------
CSS_LIGHT = """
<style>
  .stApp { background: #f6f9fc; color: #102a43; }
  .block-container { padding-top: 1.2rem; max-width: 1400px; }
  .gd-hero { background: linear-gradient(120deg, #082f49, #0369a1);
    color: white; border-radius: 18px; padding: 22px 26px; margin-bottom: 16px; }
  .gd-hero h1 { color: white; font-size: 1.8rem; margin: 0 0 6px; }
  .gd-hero p { margin: 0; color: #dbeafe; }
  .gd-card { background: white; border: 1px solid #dbe7f0; border-radius: 14px;
    padding: 14px; margin: 6px 0; color: #102a43; }
  .gd-chip { display:inline-block; border-radius:999px; padding:3px 10px; margin:2px 4px 2px 0;
    font-size:.78rem; font-weight:600; }
  .c-ok{background:#dcfce7;color:#166534;} .c-warn{background:#fef3c7;color:#92400e;}
  .c-risk{background:#fee2e2;color:#991b1b;} .c-info{background:#e0f2fe;color:#075985;}
</style>
"""
CSS_DARK = """
<style>
  .stApp { background: #0f172a; color: #e2e8f0; }
  .block-container { padding-top: 1.2rem; max-width: 1400px; }
  .gd-hero { background: linear-gradient(120deg, #020617, #0c4a6e);
    color: white; border-radius: 18px; padding: 22px 26px; margin-bottom: 16px;
    border: 1px solid #1e3a5f; }
  .gd-hero h1 { color: white; font-size: 1.8rem; margin: 0 0 6px; }
  .gd-hero p { margin: 0; color: #bae6fd; }
  .gd-card { background: #1e293b; border: 1px solid #334155; border-radius: 14px;
    padding: 14px; margin: 6px 0; color: #e2e8f0; }
  .gd-chip { display:inline-block; border-radius:999px; padding:3px 10px; margin:2px 4px 2px 0;
    font-size:.78rem; font-weight:600; }
  .c-ok{background:rgba(22,163,74,.25);color:#86efac;} .c-warn{background:rgba(245,158,11,.25);color:#fcd34d;}
  .c-risk{background:rgba(239,68,68,.25);color:#fca5a5;} .c-info{background:rgba(14,165,233,.25);color:#7dd3fc;}
  div[data-testid="stMarkdownContainer"] { color: #e2e8f0; }
</style>
"""
st.markdown(CSS_DARK if S.theme == "dark" else CSS_LIGHT, unsafe_allow_html=True)

def hero(title, sub):
    st.markdown(f'<div class="gd-hero"><h1>⛏ {title}</h1><p>{sub}</p></div>', unsafe_allow_html=True)

def chip(label, kind="info"):
    return f'<span class="gd-chip c-{kind}">{label}</span>'

# ---------- ЖУРНАЛ ----------
def log_request(uid, mode, text, meta=""):
    try:
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(json.dumps({"ts": datetime.datetime.now().isoformat(timespec="seconds"),
                                "user": uid, "mode": mode, "text": text, "meta": meta},
                               ensure_ascii=False) + "\n")
    except Exception:
        pass

def read_log():
    if not os.path.exists(LOG_FILE):
        return []
    out = []
    try:
        with open(LOG_FILE, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line:
                    try:
                        out.append(json.loads(line))
                    except Exception:
                        pass
    except Exception:
        pass
    return out

def psych_profile(entries):
    if not entries:
        return "Данных пока нет."
    texts = [e["text"].lower() for e in entries]
    joined = " ".join(texts); n = len(entries); traits = []
    if sum(t.count("?") for t in texts) / max(n, 1) >= 1.5:
        traits.append("Задаёт много уточняющих вопросов — проверяющий тип.")
    if any(w in joined for w in ["срочно", "дедлайн", "срок"]):
        traits.append("Озабочен сроками — давление дедлайнов.")
    if any(w in joined for w in ["смета", "бюджет", "стоимость"]):
        traits.append("Фокус на деньгах — оптимизация затрат.")
    if any(w in joined for w in ["фнп", "ростехнадзор", "экспертиз", "лиценз"]):
        traits.append("Осторожен в нормативной части.")
    traits.append("Кратко формулирует — решительный стиль." if sum(len(t) for t in texts)/max(n,1) < 80
                  else "Подробно описывает — вдумчивый.")
    return "\n".join(f"• {t}" for t in traits)

# ---------- ФАЙЛЫ ----------
def extract_text(file) -> str:
    name = file.name.lower()
    try:
        if name.endswith((".txt", ".md", ".csv", ".json")):
            return file.getvalue().decode("utf-8", errors="ignore")
        if name.endswith(".pdf"):
            try:
                from pypdf import PdfReader
                import io
                return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(file.getvalue())).pages[:30])
            except ImportError:
                return "[нет pypdf]"
        if name.endswith(".docx"):
            try:
                import docx, io
                d = docx.Document(io.BytesIO(file.getvalue()))
                return "\n".join(p.text for p in d.paragraphs if p.text.strip())
            except ImportError:
                return "[нет python-docx]"
        if name.endswith((".xlsx", ".xls")):
            try:
                import openpyxl, io
                wb = openpyxl.load_workbook(io.BytesIO(file.getvalue()))
                out = []
                for ws in wb.worksheets[:10]:
                    for row in ws.iter_rows(max_row=100):
                        vals = [str(c.value) for c in row if c.value is not None]
                        if vals:
                            out.append(" | ".join(vals))
                return "\n".join(out)
            except ImportError:
                return "[нет openpyxl]"
        if name.endswith(".dwg"):
            text = file.getvalue().decode("utf-16-le", errors="ignore")
            words = re.findall(r"[А-Яа-яA-Za-z0-9 .,\-()№]{6,}", text)
            return "\n".join(words[:200]) if words else "[DWG: текст не извлечён]"
    except Exception as e:
        return f"[Ошибка чтения: {e}]"
    return "[Формат не поддерживается]"

DOC_TYPES = {
    "Пояснительная записка / ПД": ["пояснительная записка", "проектная документация"],
    "ТЗ / задание на проектирование": ["задание на проектирование", "техническое задание"],
    "Инженерные изыскания": ["изыскани", "геологи", "гидрогеолог", "геодез"],
    "Уставные документы": ["устав", "бюджет", "смет", "график"],
    "Горно-техническая часть": ["горн", "вскрыти", "рудник", "шахт", "карьер"],
    "Промышленная безопасность": ["фнп", "ггэ", "экспертиз", "опасный производственный"],
    "ВОР / ведомости": ["ведомост", "вор", "объём работ", "спецификац"],
}

def classify_doc(name, text):
    blob = (name + " " + text[:5000]).lower()
    found = [t for t, keys in DOC_TYPES.items() if any(k in blob for k in keys)]
    return found or ["Прочее"]

REQUIRED_DOCS = [
    ("Пояснительная записка / ПД", "Закажите у проектировщика с СРО."),
    ("ТЗ / задание на проектирование", "Утвердите у заказчика."),
    ("Инженерные изыскания", "Геология обязательна для подземных работ."),
    ("Горно-техническая часть", "Система разработки, вскрытие, вентиляция."),
    ("Промышленная безопасность", "Раздел ПБ, ГГЭ для ОПО."),
    ("Уставные документы", "Без них проект неуправляем."),
    ("ВОР / ведомости", "Контроль объёмов и затрат."),
]

def readiness_check(present, project):
    out = [(dt, "ok", "") if dt in present else (dt, "miss", adv) for dt, adv in REQUIRED_DOCS]
    extra = []
    if project and project.get("underground"):
        extra += [("Вентиляция / ГВУ", "Проект проветривания обязателен."),
                  ("ПЛА", "Согласование с ВГСЧ до эксплуатации.")]
    if project and project.get("blasting"):
        extra.append(("Лицензия ВПХО", "Без неё работы с ВМ невозможны."))
    if project and project.get("new_project"):
        extra.append(("Обогатительная фабрика", "Отдельный проект ОФ и хвостохранилища."))
    return out, extra

# ---------- ПОИСК ----------
def search_knowledge(query, top_n=3):
    q_words = [w for w in re.split(r"\W+", query.lower()) if len(w) > 3]
    if not q_words:
        return []
    results = []
    for d in S.docs:
        text = d.get("text", "")
        if not text:
            continue
        best, best_score = "", 0
        for i in range(0, min(len(text), 100000), 500):
            ch = text[i:i+600]
            sc = sum(1 for w in q_words if w in ch.lower())
            if sc > best_score:
                best, best_score = ch, sc
        if best_score > 0:
            results.append((best_score, d["name"], best))
    results.sort(key=lambda x: -x[0])
    return results[:top_n]

# ---------- GREENFIELD ----------
def requirement_matrix(p):
    rows = [{"name": n, "group": g, "need": need, "reason": r}
            for n, g, need, r in [
        ("Устав проекта", "Управление", "Обязателен", "Цель, границы, команда, бюджет."),
        ("Календарный план / WBS", "Управление", "Обязателен", "Базовый план и критический путь."),
        ("Бюджет / финансирование", "Финансы", "Обязателен", "CAPEX/OPEX, резерв."),
        ("Реестр ИРД", "ИРД", "Обязателен", "Владелец, срок, статус."),
        ("Инженерные изыскания", "ИРД", "По составу", "ИГИ, геомеханика, гидрогеология."),
        ("Лицензия на недра", "Недропользование", "Критично", "Основание пользования недрами."),
        ("Технический проект разработки", "Проектирование", "Требуется", "Условие лицензии, ≤24 мес."),
    ]]
    def add(n, g, need, r):
        rows.append({"name": n, "group": g, "need": need, "reason": r})
    if not p.get("reserves_approved"):
        add("ТЭО кондиций и экспертиза запасов", "Недропользование", "Вероятно требуется",
            "Без сырьевой базы нельзя фиксировать мощность.")
    if p.get("capital_construction"):
        add("ПД по ПП РФ №87", "Проектирование", "Требуется", "Состав разделов по объекту.")
        add("Экспертиза ПД (ГЭЭ/НГЭ)", "Экспертизы", "Определить маршрут", "По виду объекта и финансированию.")
    if p.get("mine_method") == "подземный" or p.get("hazardous_facility"):
        add("Регистрация ОПО, ЭПБ/ГГЭ", "Промбезопасность", "Требуется", "Класс опасности и обязательства.")
    if p.get("blasting"):
        add("БВР и склад/площадка ВМ", "Промбезопасность", "Требуется", "Схема снабжения ВМ.")
    if p.get("tailings"):
        add("Хвостовое хозяйство / ГТС", "ГТС", "Требуется", "Декларация безопасности, мониторинг.")
    if p.get("water_discharge"):
        add("Водоотлив и сброс вод", "Экология/вода", "Требуется", "Разрешительные процедуры.")
    if p.get("processing"):
        add("Проект ОФ", "Проектирование", "Требуется", "Отдельный ОПО.")
    return rows

def make_roadmap(p):
    from datetime import date, timedelta
    start = p.get("target_start") or date.today()
    tasks = [
        ("M0", "Утвердить устав и базовый план", 15, [], "Инициация", "РП", True),
        ("M1", "Проверить лицензионные условия", 10, ["M0"], "Недропользование", "УЛН", True),
        ("M2", "Подтвердить запасы (ТЭО/экспертиза)", 180, ["M1"], "Недропользование", "Геология", True),
        ("M3", "Закрыть реестр ИРД", 90, ["M0"], "ИРД", "ОПИК", True),
        ("M4", "Изыскания, геомеханика, гидрогеология", 150, ["M3"], "Изыскания", "Генпроектировщик", True),
        ("M5", "ОТР и стратегия реализации", 45, ["M1", "M3"], "Концепция", "РП", True),
        ("M6", "ТЗ на техпроект", 15, ["M5"], "Проектирование", "ГИП", True),
        ("M7", "Техпроект + внутренняя экспертиза", 90, ["M2", "M4", "M6"], "Техпроект", "ГИП", True),
        ("M8", "Согласование в ЦКР (если применимо)", 45, ["M7"], "Согласования", "ГИП", True),
        ("M9", "ТЗ на ПД", 20, ["M5"], "Проектирование", "ГИП", False),
        ("M10", "ПД по №87", 120, ["M4", "M8", "M9"], "ПД", "Генпроектировщик", True),
        ("M11", "Внутренняя экспертиза ПД", 25, ["M10"], "ПД", "РП", True),
        ("M12", "Внешние экспертизы", 90, ["M11"], "Экспертизы", "УВГС", True),
        ("M13", "РД и выдача в производство", 110, ["M10"], "РД", "Генпроектировщик", False),
        ("M14", "Контрактация и long-lead", 120, ["M5"], "Закупки", "МТО", False),
        ("M15", "Подготовка площадки", 45, ["M12", "M13"], "Строительство", "ЗУД", True),
        ("M16", "СМР / ГКР", 365, ["M15", "M14"], "Строительство", "ЗУД", True),
        ("M17", "ПНР и ввод", 60, ["M16"], "Ввод", "Заказчик", True),
    ]
    if not p.get("capital_construction"):
        skip = {"M10", "M11", "M12", "M13", "M15", "M16", "M17"}
        tasks = [t for t in tasks if t[0] not in skip]
    end_by = {}; result = []
    for tid, work, dur, preds, phase, owner, crit in tasks:
        pe = max((end_by[x] for x in preds if x in end_by), default=start - timedelta(days=1))
        ts_ = pe + timedelta(days=1); te = ts_ + timedelta(days=dur - 1)
        end_by[tid] = te
        result.append({"id": tid, "work": work, "phase": phase, "owner": owner,
                       "start": ts_, "finish": te, "duration": dur,
                       "predecessors": preds, "critical": crit})
    return result

# ---------- ШТАБ: синхронизация ----------
def sync_all():
    notes = []
    if S.gf_project and not S.hq_facts:
        gf = S.gf_project
        S.hq_facts.update({
            "project_name": gf.get("project_name", ""),
            "customer": gf.get("company", ""),
            "location": gf.get("region", ""),
        })
        notes.append("Greenfield: профиль проекта импортирован.")
    if S.gf_tasks and not S.roadmap:
        for t in S.gf_tasks:
            S.roadmap.append({"work": t["work"], "start": str(t["start"]),
                              "finish": str(t["finish"]), "pct": 0,
                              "criterion": f"Владелец: {t['owner']}"})
        notes.append(f"Greenfield: {len(S.gf_tasks)} работ в дорожную карту.")
    if S.project and not S.hq_facts.get("project_name"):
        S.hq_facts["project_name"] = S.project["name"]
        notes.append("Мастер: название проекта.")
    return notes or ["Новых данных нет — пройдите диагностику или мастер."]

# ---------- ГЕНЕРАТОРЫ УСТАВА И ТЗ ----------
def txt(v, empty="Требует заполнения"):
    v = (v or "").strip() if isinstance(v, str) else v
    return v if v else empty

def generate_charter(f):
    tl = "\n".join(f"| {r['work']} | {r['start']} | {r['finish']} | {r['criterion']} |" for r in S.roadmap) or "| — | — | — | — |"
    bl = "\n".join(f"| {b.get('item','')} | {b.get('y1',0)} | {b.get('y2',0)} | {b.get('total',0)} |" for b in S.budget) or "| — | — | — | — |"
    rl = "\n".join(f"| {r.get('risk','')} | {r.get('prob','')} | {r.get('owner','')} |" for r in S.risks) or "| — | — | — |"
    return f"""# УСТАВ ПРОЕКТА «{txt(f.get('project_name'))}»
**РП:** {txt(f.get('project_manager'))} · **Заказчик:** {txt(f.get('customer'))}

## 1. Основание и даты
Основание: {txt(f.get('basis'))} · Завершение: {txt(f.get('close_date'))}

## 2. Границы
Место: {txt(f.get('location'))} · Лицензия: {txt(f.get('license'))}
{txt(f.get('description'))}

## 3. Цель и результаты
Цель: {txt(f.get('goal'))}

## 4. Этапы
| Работа | Начало | Окончание | Критерий |
|---|---|---|---|
{tl}

## 5. Бюджет (тыс. ₽)
| Статья | Год 1 | Год 2 | Итого |
|---|---|---|---|
{bl}

## 6. Риски
| Риск | Вероятность | Владелец |
|---|---|---|
{rl}

> Черновик. Проверяется ответственными лицами; не заменяет экспертизу."""

def generate_tz(f):
    return f"""# ТЗ «{txt(f.get('tz_title'))}»
**Тип:** {f.get('tz_type', 'Проектная документация')} · **Заказчик:** {txt(f.get('customer'))}

| № | Данные | Содержание |
|---:|---|---|
| 1 | Объект, место | {txt(f.get('tz_title'))}; {txt(f.get('location'))} |
| 2 | Основание | {txt(f.get('basis'))} |
| 3 | Лицензия/права | {txt(f.get('license'))} |
| 4 | Финансирование | {txt(f.get('funding_source'))} |
| 5 | ТЭП | {txt(f.get('tep'))} |
| 6 | Технологии | {txt(f.get('tech'))} |
| 7 | Инженерное обеспечение | {txt(f.get('eng'))} |
| 8 | Исходные данные | {txt(f.get('source_data'))} |
| 9 | Экспертизы | {txt(f.get('expertise'))} |

## Комплектность ИРД
- [ ] Лицензия на недра
- [ ] Права на землю
- [ ] Изыскания
- [ ] Запасы/кондиции

> Полнота ИРД подтверждается проектировщиком и заказчиком."""

# ============================================================
#  ИНТЕРФЕЙС
# ============================================================
st.title("⛏ Горный Дельта")
st.caption("ИИ-платформа руководителя проектов • v3.0 (единый файл)")

user_id = st.sidebar.text_input("Ваш ID / e-mail", value="guest")

# ---------- ПЕРЕКЛЮЧАТЕЛЬ ТЕМЫ ----------
S.theme = st.sidebar.radio("Тема", ["light", "dark"], format_func=lambda t: "☀️ Светлая" if t == "light" else "🌙 Тёмная",
                           index=0 if S.theme == "light" else 1, key="theme_radio")

menu = st.sidebar.radio("Навигация", [
    "🏗️ Цифровой штаб", "🗺 Greenfield-диагностика", "🚀 Новый проект",
    "📂 Документы", "✅ Чек-лист", "🧠 Инженерная логика",
    "💬 Ассистент", "📐 PRO: DWG", "💳 Подписка", "🔐 Владелец"])

# ---------- ШТАБ ----------
if menu.startswith("🏗️"):
    hero("Цифровой штаб", "Паспорт, дорожная карта, бюджет, риски, Устав и ТЗ — всё в одном месте.")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Работ в карте", len(S.roadmap))
    c2.metric("Документов", len(S.docs))
    c3.metric("Требований", len(S.gf_requirements))
    c4.metric("Greenfield", "✓" if S.gf_project else "—")
    if st.button("🔄 Синхронизировать", type="primary"):
        for n in sync_all():
            st.markdown(f'<div class="gd-card">✓ {n}</div>', unsafe_allow_html=True)
    f = S.hq_facts
    page = st.radio("Раздел", ["Паспорт", "Дорожная карта", "Бюджет и риски", "Устав", "ТЗ"],
                    horizontal=True)
    if page == "Паспорт":
        c1, c2 = st.columns(2)
        f["project_name"] = c1.text_input("Проект", f.get("project_name", ""))
        f["customer"] = c2.text_input("Заказчик", f.get("customer", ""))
        f["project_manager"] = c1.text_input("РП", f.get("project_manager", ""))
        f["location"] = c2.text_input("Местоположение", f.get("location", ""))
        f["license"] = c1.text_input("Лицензия", f.get("license", ""))
        f["close_date"] = c2.text_input("Целевая дата", f.get("close_date", ""))
        f["basis"] = st.text_area("Основание", f.get("basis", ""), height=60)
        f["goal"] = st.text_area("Цель", f.get("goal", ""), height=60)
        f["description"] = st.text_area("Границы", f.get("description", ""), height=80)
        if S.docs:
            st.subheader("Доказательная база (база знаний)")
            q = st.text_input("Поиск по документам", placeholder="лицензия, заказчик...")
            if q:
                for sc, fn_, frag in search_knowledge(q, 2):
                    st.info(f"**{fn_}:** {frag[:300]}")
    elif page == "Дорожная карта":
        if not S.roadmap:
            st.info("Карта пуста — синхронизируйте (кнопка сверху) или добавьте работы.")
        for i, r in enumerate(S.roadmap):
            c1, c2, c3, c4, c5 = st.columns([4, 2, 2, 1, 2])
            r["work"] = c1.text_input("Работа", r["work"], key=f"w{i}")
            r["start"] = c2.text_input("Начало", str(r["start"]), key=f"s{i}")
            r["finish"] = c3.text_input("Окончание", str(r["finish"]), key=f"f{i}")
            r["pct"] = c4.number_input("%", 0, 100, int(r["pct"]), key=f"p{i}")
            r["criterion"] = c5.text_input("Критерий", r["criterion"], key=f"c{i}")
        if st.button("➕ Добавить работу"):
            S.roadmap.append({"work": "Новая работа", "start": "", "finish": "", "pct": 0, "criterion": ""})
            st.rerun()
    elif page == "Бюджет и риски":
        left, right = st.columns(2)
        with left:
            st.subheader("Бюджет, тыс. ₽")
            for i, b in enumerate(S.budget):
                b["item"] = st.text_input("Статья", b.get("item", ""), key=f"bi{i}")
                b["y1"] = st.number_input(f"Год 1 ({b.get('item','')})", 0, None, int(b.get("y1", 0)), key=f"by1_{i}")
                b["y2"] = st.number_input(f"Год 2 ({b.get('item','')})", 0, None, int(b.get("y2", 0)), key=f"by2_{i}")
                b["total"] = b["y1"] + b["y2"]
            if st.button("➕ Статья"):
                S.budget.append({"item": "", "y1": 0, "y2": 0, "total": 0}); st.rerun()
            st.metric("ИТОГО", sum(b.get("total", 0) for b in S.budget))
        with right:
            st.subheader("Риски")
            for i, r in enumerate(S.risks):
                r["risk"] = st.text_input("Риск", r.get("risk", ""), key=f"ri{i}")
                probs = ["Низкая", "Средняя", "Высокая"]
                r["prob"] = st.selectbox("Вероятность", probs,
                                         index=probs.index(r["prob"]) if r.get("prob") in probs else 1, key=f"rp{i}")
                r["owner"] = st.text_input("Владелец", r.get("owner", ""), key=f"ro{i}")
            if st.button("➕ Риск"):
                S.risks.append({"risk": "", "prob": "Средняя", "owner": ""}); st.rerun()
    elif page == "Устав":
        ch = generate_charter(S.hq_facts)
        st.markdown(ch)
        st.download_button("⬇️ Скачать Устав (.md)", ch, "ustav_draft.md", "text/markdown")
    else:
        f["tz_title"] = st.text_input("Объект ТЗ", f.get("tz_title", f.get("project_name", "")))
        f["tz_type"] = st.text_input("Тип ТЗ", f.get("tz_type", "Проектная документация"))
        f["tep"] = st.text_area("ТЭП", f.get("tep", ""), height=60)
        f["tech"] = st.text_area("Технологии", f.get("tech", ""), height=80)
        tz = generate_tz(f)
        st.markdown(tz)
        st.download_button("⬇️ Скачать ТЗ (.md)", tz, "tz_draft.md", "text/markdown")

# ---------- GREENFIELD ----------
elif menu.startswith("🗺"):
    hero("Greenfield-диагностика", "От лицензии и запасов до ввода в эксплуатацию.")
    if not S.gf_project:
        with st.form("gf"):
            a, b = st.columns(2)
            with a:
                pn = a.text_input("Название проекта *")
                comp = b.text_input("Заказчик *")
                reg = a.text_input("Регион")
                method = st.radio("Способ разработки", ["подземный", "открытый", "комбинированный", "неизвестно"], horizontal=True)
                stage = st.selectbox("Этап", ["вскрытие", "добыча", "разведка", "инфраструктура"])
                tstart = st.date_input("Целевая дата старта")
            with b:
                has_lic = st.radio("Лицензия", ["есть", "нет", "неизвестно"], horizontal=True)
                reserves = st.radio("Запасы утверждены?", ["да", "нет", "неизвестно"], horizontal=True)
                capital = st.radio("Объекты капстроительства?", ["да", "нет", "неизвестно"], horizontal=True)
                blasting = st.checkbox("БВР / ВМ", True)
                tailings = st.checkbox("Хвостовое хозяйство / ГТС")
                water = st.checkbox("Водоотлив / сброс")
                processing = st.checkbox("Фабрика / переработка")
                hazardous = st.checkbox("ОПО", True)
            if st.form_submit_button("Построить дорожную карту", type="primary"):
                S.gf_project = {"project_name": pn or "Greenfield", "company": comp, "region": reg,
                                "mine_method": method, "stage": stage, "target_start": tstart,
                                "has_license": has_lic == "есть", "reserves_approved": reserves == "да",
                                "capital_construction": capital == "да", "blasting": blasting,
                                "tailings": tailings, "water_discharge": water,
                                "processing": processing, "hazardous_facility": hazardous}
                S.gf_tasks = make_roadmap(S.gf_project)
                S.gf_requirements = requirement_matrix(S.gf_project)
                log_request(user_id, "greenfield", json.dumps(S.gf_project, ensure_ascii=False, default=str))
                st.rerun()
    else:
        p = S.gf_project
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Работ", len(S.gf_tasks))
        c2.metric("Критических", sum(1 for t in S.gf_tasks if t["critical"]))
        c3.metric("Требований", len(S.gf_requirements))
        c4.metric("Прогноз завершения", max(t["finish"] for t in S.gf_tasks).strftime("%d.%m.%Y") if S.gf_tasks else "—")
        blockers = []
        if not p.get("has_license"):
            blockers.append("Подтвердить лицензию и её условия")
        if not p.get("reserves_approved"):
            blockers.append("Подтвердить запасы / маршрут утверждения")
        if p.get("mine_method") == "подземный" and not p.get("water_discharge"):
            blockers.append("Решения по водоотливу и водоотведению")
        st.subheader("Следующие действия")
        for it in blockers or ["Блокеров не выявлено — переходите к ИРД и ТЗ."]:
            st.markdown(f'<div class="gd-card"><b>Приоритет:</b> {it}</div>', unsafe_allow_html=True)
        tab1, tab2 = st.tabs(["🗺 Дорожная карта", "📚 Документация"])
        with tab1:
            phases = list(dict.fromkeys(t["phase"] for t in S.gf_tasks))
            for ph in phases:
                with st.expander(ph, expanded=ph in ("Инициация", "Недропользование")):
                    for t in [x for x in S.gf_tasks if x["phase"] == ph]:
                        mark = "🔴" if t["critical"] else "🔵"
                        st.markdown(f"{mark} **{t['id']} — {t['work']}**")
                        st.caption(f"{t['start']:%d.%m.%Y} → {t['finish']:%d.%m.%Y} · {t['duration']} дн. · {t['owner']}")
        with tab2:
            for row in S.gf_requirements:
                kind = "risk" if row["need"] == "Критично" else "warn" if "Требуется" in row["need"] else "info"
                with st.expander(f"{row['group']} · {row['name']} · {row['need']}"):
                    st.markdown(chip(row["need"], kind), unsafe_allow_html=True)
                    st.write(row["reason"])
        if st.button("🔄 Перепройти диагностику"):
            S.gf_project, S.gf_tasks, S.gf_requirements = {}, [], []
            st.rerun()

# ---------- МАСТЕР ----------
elif menu.startswith("🚀"):
    hero("Новый проект", "Шаг 1 — базовые параметры")
    with st.form("wizard"):
        name = st.text_input("Название проекта", "")
        c1, c2 = st.columns(2)
        with c1:
            underground = st.checkbox("Подземные работы")
            open_pit = st.checkbox("Открытые работы")
            blasting = st.checkbox("БВР", True)
        with c2:
            infrastructure = st.checkbox("Инфраструктура есть")
            has_files = st.checkbox("Есть исходные файлы")
            new_project = st.checkbox("Проект новый")
        volume = st.number_input("Добыча, т/год", 0, None, 500000, 50000)
        if st.form_submit_button("➡️ Создать проект", type="primary") and name:
            S.project = {"name": name, "underground": underground, "open_pit": open_pit,
                         "blasting": blasting, "infrastructure": infrastructure,
                         "has_files": has_files, "new_project": new_project, "volume": volume}
            S.wizard_done = True
            log_request(user_id, "wizard", name)
            st.success(f"Проект «{name}» создан!")
    if S.wizard_done:
        st.info(f"Текущий проект: «{S.project['name']}»")
        if st.button("🔄 Сбросить проект"):
            S.wizard_done, S.docs, S.project = False, [], None
            st.rerun()

# ---------- ДОКУМЕНТЫ ----------
elif menu.startswith("📂"):
    hero("Документы проекта", "Загружайте ПД, ТЗ, изыскания, ВОР, DWG — система учится на них.")
    files = st.file_uploader("Файлы", type=["pdf", "docx", "txt", "md", "csv", "xlsx", "dwg"],
                             accept_multiple_files=True)
    if files:
        for f_ in files:
            text = extract_text(f_)
            types = classify_doc(f_.name, text)
            try:
                with open(os.path.join(DOCS_DIR, f"{datetime.datetime.now():%Y%m%d%H%M%S}_{f_.name}"), "wb") as o:
                    o.write(f_.getvalue())
            except Exception:
                pass
            S.docs.append({"name": f_.name, "text": text, "types": types})
            log_request(user_id, "upload", f_.name, ",".join(types))
        st.success(f"Загружено: {len(files)}")
    for d in S.docs:
        with st.expander(f"📄 {d['name']} — {', '.join(d['types'])}"):
            st.text(d["text"][:1200] or "[текст не извлечён]")

# ---------- ЧЕК-ЛИСТ ----------
elif menu.startswith("✅"):
    hero("Чек-лист готовности", "Чего не хватает и что делать")
    present = set()
    for d in S.docs:
        present.update(d["types"])
    base, extra = readiness_check(present, S.project)
    ok = sum(1 for _, s_, _ in base if s_ == "ok")
    st.progress(ok / len(base), text=f"Готовность: {ok} из {len(base)}")
    for dt, st_, adv in base:
        if st_ == "ok":
            st.markdown(f"✅ **{dt}** — загружено")
        else:
            st.markdown(f"⚠️ **{dt}** — отсутствует → {adv}")
    for dt, adv in extra:
        st.markdown(f"🔴 **{dt}** — требуется → {adv}")

# ---------- ЛОГИКА ----------
elif menu.startswith("🧠"):
    hero("Инженерная логика", "Цепочка обязательных требований")
    if not S.project:
        st.info("Сначала создайте проект в «🚀 Новый проект».")
    elif st.button("Построить цепочку требований", type="primary"):
        log_request(user_id, "logic", json.dumps(S.project, ensure_ascii=False))
        reqs = []
        if S.project.get("underground"):
            reqs += ["Подземные работы → ОПО I класса", "Регистрация ОПО в Ростехнадзоре",
                     "Техпроект разработки ≤24 мес.", "Экспертиза ГГЭ — обязательна",
                     "Вентиляция/ГВУ, паспорта крепления, ПЛА", "БВР → лицензия ВПХО"]
        if S.project.get("open_pit"):
            reqs += ["Открытые работы → класс ОПО по добыче", "ПГР с согласованием",
                     "Паспорт БВР", "ГГЭ для объектов под ФНП"]
        for i, r in enumerate(reqs, 1):
            st.markdown(f'{i}. {r}')
        st.info("💡 Инфраструктуру (ГСМ, РММ, АБК) выносите в отдельный проект под негосэкспертизу.")

# ---------- АССИСТЕНТ ----------
elif menu.startswith("💬"):
    hero("Ассистент по базе знаний", "Ответы строятся только на ваших документах")
    if not S.docs:
        st.info("База пуста — загрузите документы в разделе «📂 Документы».")
    else:
        q = st.text_input("Вопрос", placeholder="какая система разработки принята?")
        if st.button("Спросить") and q.strip():
            log_request(user_id, "ask", q)
            res = search_knowledge(q)
            if res:
                for sc, fn_, frag in res:
                    st.markdown(f"**📄 {fn_}:**")
                    st.info(frag[:600])
            else:
                st.warning("Не найдено. Уточните формулировку или загрузите документы.")

# ---------- DWG ----------
elif menu.startswith("📐"):
    hero("PRO: анализ DWG-чертежей", f"Тариф PRO — {PRICE_PRO} ₽/мес")
    if S.plan != "pro":
        st.warning("🔒 Активируйте PRO в разделе «💳 Подписка».")
    else:
        f_ = st.file_uploader("DWG-файл", type=["dwg"])
        if f_ and st.button("📐 Проанализировать"):
            log_request(user_id, "dwg", f_.name)
            text = extract_text(f_).lower()
            st.markdown("**Извлечено:**")
            st.caption(text[:600] or "[пусто]")
            recs = []
            if "лоток" in text: recs.append("🔧 Лотки: толщина 2,5 → 1,5 мм (экономия до 40% металла).")
            if "кабел" in text: recs.append("🔧 Кабели: проверить уменьшение сечения на 1 ступень.")
            if "металлоконструкц" in text: recs.append("🔧 МК: типовые серии вместо индивидуальных (−10–20%).")
            if "трубопровод" in text: recs.append("🔧 Трубопроводы: пересчёт толщины стенки по износу.")
            st.markdown("**Рекомендации:**")
            st.markdown("\n\n".join(recs) or "Текст из DWG ограничен — нужен ИИ-модуль распознавания.")

# ---------- ПОДПИСКА ----------
elif menu.startswith("💳"):
    hero("Подписка", "Базовая и PRO")
    st.markdown(f"**Базовая** — {PRICE_BASIC} ₽/мес: мастер, документы, чек-лист, штаб, ассистент.")
    st.markdown(f"**PRO** — {PRICE_PRO} ₽/мес: + анализ DWG-чертежей.")
    promo = st.text_input("Промокод").strip().upper()
    if st.button("Активировать"):
        log_request(user_id, "billing", promo)
        if promo in PROMO_CODES:
            S.plan = PROMO_CODES[promo]
            st.success(f"Тариф активирован: {S.plan.upper()}")
        else:
            st.error("Неверный промокод.")

# ---------- ВЛАДЕЛЕЦ ----------
elif menu.startswith("🔐"):
    hero("Кабинет владельца", "Журнал, портреты пользователей, база знаний")
    if not S.show_owner:
        pwd = st.text_input("Пароль владельца", type="password")
        if st.button("Войти") and pwd == OWNER_PASSWORD:
            S.show_owner = True
            st.rerun()
    if S.show_owner:
        entries = read_log()
        st.subheader(f"Журнал: {len(entries)} записей")
        users = {}
        for e in entries:
            users.setdefault(e["user"], []).append(e)
        for u, es in users.items():
            with st.expander(f"{u} — {len(es)} действий"):
                st.markdown("**Психологический портрет:**")
                st.markdown(psych_profile(es))
                st.markdown("**История:**")
                for e in es[-20:]:
                    st.text(f"[{e['ts']}] ({e['mode']}) {e['text'][:120]}")
        st.subheader("База знаний на сервере")
        for fn_ in sorted(os.listdir(DOCS_DIR)):
            st.text(fn_)

st.sidebar.markdown("---")
st.sidebar.caption("© 2026 Горный Дельта v3.0")
