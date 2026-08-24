# ============================================================
#  ГОРНЫЙ ДЕЛЬТА v2.2 — ИИ-платформа руководителя проектов
# ============================================================
#  ИЗМЕНЕНИЯ v2.2:
#   - Добавлен «Цифровой штаб» (project_hq.py) — агрегатор:
#     синхронизируется с Greenfield, AI Copilot, мастером,
#     базой знаний; формирует Устав и ТЗ.
#   - Все импорты модулей защищены try/except.
#  ФАЙЛЫ В РЕПОЗИТОРИИ (корень):
#    app.py, project_hq.py, greenfield_command_center.py,
#    ai_project_copilot.py (опционально), requirements.txt
#  requirements.txt:
#    streamlit>=1.36
#    pypdf
#    python-docx
#    openpyxl
#    pandas
# ============================================================

import streamlit as st
import json
import os
import datetime
import re

# ---------- БЕЗОПАСНЫЙ ИМПОРТ МОДУЛЕЙ ----------
GREENFIELD_OK, GREENFIELD_ERROR = False, ""
try:
    from greenfield_command_center import apply_command_center_theme, greenfield_page
    GREENFIELD_OK = True
except Exception as e:
    GREENFIELD_ERROR = str(e)

HQ_OK, HQ_ERROR = False, ""
try:
    from project_hq import project_hq_page
    HQ_OK = True
except Exception as e:
    HQ_ERROR = str(e)

COPILOT_OK, COPILOT_ERROR = False, ""
try:
    from ai_project_copilot import render_ai_copilot
    COPILOT_OK = True
except Exception as e:
    COPILOT_ERROR = str(e)

# ---------- НАСТРОЙКИ ВЛАДЕЛЬЦА ----------
OWNER_PASSWORD = "CHANGE_ME_2026"
PRICE_BASIC = 4900
PRICE_PRO = 14900
PROMO_CODES = {"PILOT2026": "pro", "DEMO": "basic"}
LOG_FILE = "users_log.jsonl"
DOCS_DIR = "knowledge"
os.makedirs(DOCS_DIR, exist_ok=True)

st.set_page_config(page_title="Горный Дельта — ИИ-платформа РП", page_icon="⛏", layout="wide")
if GREENFIELD_OK:
    apply_command_center_theme()

# ---------- СОСТОЯНИЕ ----------
S = st.session_state
defaults = {"plan": "free", "show_owner": False, "wizard_done": False,
            "project": None, "docs": [], "messages": []}
for k, v in defaults.items():
    if k not in S:
        S[k] = v

# ---------- ЖУРНАЛ ----------
def log_request(user_id, mode, text, meta=""):
    entry = {"ts": datetime.datetime.now().isoformat(timespec="seconds"),
             "user": user_id, "mode": mode, "text": text, "meta": meta}
    try:
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")
    except Exception:
        pass

def read_log():
    if not os.path.exists(LOG_FILE):
        return []
    out = []
    try:
        with open(LOG_FILE, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line:
                    try:
                        out.append(json.loads(line))
                    except json.JSONDecodeError:
                        pass
    except Exception:
        pass
    return out

# ---------- ПСИХОЛОГИЧЕСКИЙ ПОРТРЕТ ----------
def psych_profile(entries):
    if not entries:
        return "Данных пока нет."
    texts = [e["text"].lower() for e in entries]
    joined = " ".join(texts)
    n = len(entries)
    traits = []
    if sum(t.count("?") for t in texts) / max(n, 1) >= 1.5:
        traits.append("Задаёт много уточняющих вопросов — аккуратный, проверяющий тип.")
    if any(w in joined for w in ["срочно", "быстрее", "дедлайн", "срок"]):
        traits.append("Озабочен сроками — работает под давлением дедлайнов.")
    if any(w in joined for w in ["смета", "бюджет", "стоимость", "экономия"]):
        traits.append("Фокус на деньгах — прагматик, ищет оптимизацию затрат.")
    if any(w in joined for w in ["норм", "фнп", "ростехнадзор", "экспертиз", "лиценз"]):
        traits.append("Осторожен в нормативной части.")
    if any(w in joined for w in ["письмо", "протокол", "повестк", "устав"]):
        traits.append("Делегирует рутину — бережёт время для управления.")
    avg_len = sum(len(t) for t in texts) / max(n, 1)
    traits.append("Кратко формулирует — решительный стиль." if avg_len < 80
                  else "Подробно описывает задачи — вдумчивый, любит контекст.")
    return "\n".join(f"• {t}" for t in traits)

# ---------- ИЗВЛЕЧЕНИЕ ТЕКСТА ----------
def extract_text(file) -> str:
    name = file.name.lower()
    try:
        if name.endswith((".txt", ".md", ".csv", ".json")):
            return file.getvalue().decode("utf-8", errors="ignore")
        if name.endswith(".pdf"):
            try:
                from pypdf import PdfReader
                import io
                reader = PdfReader(io.BytesIO(file.getvalue()))
                return "\n".join((p.extract_text() or "") for p in reader.pages[:30])
            except ImportError:
                return "[pip install pypdf]"
        if name.endswith(".docx"):
            try:
                import docx
                import io
                d = docx.Document(io.BytesIO(file.getvalue()))
                return "\n".join(p.text for p in d.paragraphs if p.text.strip())
            except ImportError:
                return "[pip install python-docx]"
        if name.endswith((".xlsx", ".xls")):
            try:
                import openpyxl
                import io
                wb = openpyxl.load_workbook(io.BytesIO(file.getvalue()))
                out = []
                for ws in wb.worksheets[:10]:
                    for row in ws.iter_rows(max_row=100):
                        vals = [str(c.value) for c in row if c.value is not None]
                        if vals:
                            out.append(" | ".join(vals))
                return "\n".join(out)
            except ImportError:
                return "[pip install openpyxl]"
        if name.endswith(".dwg"):
            raw = file.getvalue()
            try:
                text = raw.decode("utf-16-le", errors="ignore")
            except Exception:
                text = ""
            words = re.findall(r"[А-Яа-яA-Za-z0-9 .,\-()№]{6,}", text)
            return "\n".join(words[:200]) if words else "[DWG: текст не извлечён]"
    except Exception as e:
        return f"[Ошибка чтения файла: {e}]"
    return "[Формат не поддерживается]"

# ---------- КЛАССИФИКАЦИЯ ----------
DOC_TYPES = {
    "Пояснительная записка / ПД": ["пояснительная записка", "пз", "проектная документация"],
    "ТЗ / задание на проектирование": ["задание на проектирование", "тз", "техническое задание"],
    "Инженерные изыскания": ["изыскани", "геологи", "гидрогеолог", "геодез", "экологическ"],
    "Уставные документы": ["устав", "бюджет", "смет", "график", "календарн"],
    "Горно-техническая часть": ["горн", "вскрыти", "система разработк", "рудник", "шахт", "ствол", "карьер"],
    "Промышленная безопасность": ["безопасн", "фнп", "ггэ", "экспертиз", "опасный производственный", "опо"],
    "ВОР / ведомости": ["ведомост", "вор", "объём работ", "объем работ", "спецификац"],
}

def classify_doc(name, text):
    blob = (name + " " + text[:5000]).lower()
    found = [t for t, keys in DOC_TYPES.items() if any(k in blob for k in keys)]
    return found or ["Прочее"]

# ---------- ЧЕК-ЛИСТ ----------
REQUIRED_DOCS = [
    ("Пояснительная записка / ПД", "Основной документ проекта. Если нет — закажите у проектировщика с СРО."),
    ("ТЗ / задание на проектирование", "Без ТЗ проектировщик не начнёт. Утвердите у заказчика."),
    ("Инженерные изыскания", "Геология/гидрогеология обязательны для подземных работ."),
    ("Горно-техническая часть", "Система разработки, вскрытие, вентиляция — ядро проекта рудника."),
    ("Промышленная безопасность", "Раздел ПБ, ГГЭ — обязательно для ОПО."),
    ("Уставные документы", "Устав, бюджет, график — без них проект неуправляем."),
    ("ВОР / ведомости", "Нужны для контроля объёмов и оптимизации затрат."),
]

def readiness_check(doc_types_present, project):
    out = []
    for dt, advice in REQUIRED_DOCS:
        out.append((dt, "ok", "") if dt in doc_types_present else (dt, "miss", advice))
    extra = []
    if project.get("underground"):
        extra.append(("Проект вентиляции / ГВУ", "miss", "Для подземных работ обязателен проект проветривания."))
        extra.append(("ПЛА", "miss", "Составляется до начала эксплуатации, согласуется с ВГСЧ."))
    if project.get("blasting"):
        extra.append(("Решение по ВМ + лицензия ВПХО", "miss", "Без лицензии ВПХО работы с ВМ невозможны."))
    if not project.get("infrastructure"):
        extra.append(("Инфраструктура (АБК, ГСМ, РММ)", "miss", "Вынесите в отдельный проект под негосэкспертизу."))
    if project.get("new_project"):
        extra.append(("Обогатительная фабрика", "miss", "Требуется отдельный проект ОФ и хвостовго хозяйства."))
    return out, extra

# ---------- ПОИСК ПО БАЗЕ ЗНАНИЙ ----------
def search_knowledge(query, top_n=3):
    q_words = [w for w in re.split(r"\W+", query.lower()) if len(w) > 3]
    if not q_words:
        return []
    results = []
    for d in S.docs:
        text = d.get("text", "")
        if not text:
            continue
        chunks = [text[i:i+600] for i in range(0, min(len(text), 100000), 500)]
        best, best_score = "", 0
        for ch in chunks:
            score = sum(1 for w in q_words if w in ch.lower())
            if score > best_score:
                best, best_score = ch, score
        if best_score > 0:
            results.append((best_score, d["name"], best))
    results.sort(key=lambda x: -x[0])
    return results[:top_n]

# ---------- ИНЖЕНЕРНАЯ ЛОГИКА ----------
def build_requirements(p):
    reqs = []
    if p.get("underground"):
        reqs += [
            "Подземные работы → ОПО (как правило, I класса)",
            "Регистрация ОПО в реестре Ростехнадзора",
            "Технический проект разработки (условие лицензии, ≤24 мес.)",
            "Экспертиза ГГЭ — обязательна",
            "Проект вентиляции / ГВУ, паспорта крепления, ПЛА",
            "Если БВР → лицензия ВПХО",
        ]
    if p.get("open_pit"):
        reqs += [
            "Открытые работы → классификация ОПО по объёму добычи",
            "ПГР с согласованием Ростехнадзора",
            "При БВР → лицензия ВПХО, паспорт БВР",
            "ГГЭ для объектов под ФНП",
        ]
    return reqs

# ---------- DWG ----------
def analyze_dwg(text, name):
    recs = []
    t = text.lower()
    if "лоток" in t:
        recs.append("🔧 Кабельные лотки: толщина 2,5 → 1,5 мм на непроходных участках (экономия до 40%).")
    if "кабел" in t:
        recs.append("🔧 Сечения кабелей: проверьте уменьшение на 1 ступень.")
    if "металлоконструкц" in t:
        recs.append("🔧 Металлоконструкции: типовые решения — экономия 10–20%.")
    if "трубопровод" in t:
        recs.append("🔧 Трубопроводы: пересчитайте толщину стенки по износу.")
    if not recs:
        recs.append("Извлечённый из DWG текст ограничен. Полный анализ — при подключении ИИ-модуля.")
    return "\n".join(recs)

# ---------- ГЕНЕРАТОРЫ ----------
def gen_charter(p):
    return ("УСТАВ ПРОЕКТА\n"
            f"1. Название: {p.get('name')}\n"
            f"2. Цель: {'вскрытие и отработка месторождения' if p.get('new_project') else 'реализация проекта'}\n"
            "3. Границы: от ИРД до ввода в эксплуатацию\n"
            "4. Критерии успеха: заключения ГГЭ/ГЭЭ, ввод в мощность, бюджет ±10%\n"
            "5. Роль РП: координация подрядчиков, сроки, бюджет, надзор\n"
            "6. Ограничения: ФНП, лицензионные условия, экология\n")

def gen_schedule(p):
    stages = [("ИРД и лицензирование", 3), ("ТЗ и изыскания", 4), ("Техпроект / ПД", 6),
              ("Экспертизы ГГЭ/ГЭЭ", 4), ("РД", 6), ("Строительство", 12), ("ПНР и ввод", 3)]
    lines = ["ГРАФИК ПРОЕКТА (месяцы):"]
    t = 1
    for n, d in stages:
        lines.append(f"  Мес. {t:>2}–{t+d-1:>2}: {n}")
        t += d
    lines.append(f"  ИТОГО: ~{t-1} мес.")
    return "\n".join(lines)

def gen_budget(p):
    return ("БЮДЖЕТ (укрупнённо):\n  • Проектирование и экспертизы\n  • Горно-капитальные работы\n"
            "  • Оборудование\n  • Инфраструктура (отдельный проект)\n  • Резерв 10%\n"
            "⚠️ Смета не разрабатывалась — приоритетная задача.")

# ============================================================
#  ИНТЕРФЕЙС
# ============================================================
st.title("⛏ Горный Дельта")
st.caption("ИИ-платформа руководителя проектов горнодобывающего предприятия • v2.2")

user_id = st.sidebar.text_input("Ваш ID / e-mail", value="guest")

menu = st.sidebar.radio("Навигация", [
    "🏗️ Цифровой штаб",
    "🗺 Greenfield command center",
    "🤖 AI Project Copilot",
    "🚀 Новый проект (мастер)",
    "📂 Документы проекта",
    "✅ Чек-лист готовности",
    "🧠 Инженерная логика",
    "💬 Ассистент (база знаний)",
    "📋 Уставные документы",
    "📐 PRO: анализ DWG",
    "💳 Подписка",
    "🔐 Кабинет владельца",
])

# ---------- ЦИФРОВОЙ ШТАБ ----------
if menu.startswith("🏗️"):
    if HQ_OK:
        project_hq_page(S)
    else:
        st.error("Модуль «Цифровой штаб» не загружен.")
        st.code(HQ_ERROR, language="text")
        st.info("Проверьте, что project_hq.py лежит в корне репозитория рядом с app.py "
                "и что в requirements.txt добавлен pandas.")

# ---------- GREENFIELD ----------
elif menu.startswith("🗺"):
    if GREENFIELD_OK:
        greenfield_page(S)
    else:
        st.error("Модуль Greenfield не загружен.")
        st.code(GREENFIELD_ERROR, language="text")

# ---------- AI COPILOT ----------
elif menu.startswith("🤖"):
    if COPILOT_OK:
        render_ai_copilot(
            state=S,
            docs=S.docs,
            search_fn=search_knowledge,
            greenfield_project=S.gf_project if "gf_project" in S else {},
        )
    else:
        st.error("AI Project Copilot не загружен: " + COPILOT_ERROR)
        st.info("Положите корректный файл ai_project_copilot.py в корень репозитория.")

# ---------- МАСТЕР ----------
elif menu.startswith("🚀"):
    st.header("Создание проекта — Шаг 1: параметры")
    with st.form("wizard"):
        name = st.text_input("Название проекта", "")
        c1, c2 = st.columns(2)
        with c1:
            underground = st.checkbox("Подземные горные работы")
            open_pit = st.checkbox("Открытые горные работы")
            blasting = st.checkbox("Применяются БВР", True)
        with c2:
            infrastructure = st.checkbox("Инфраструктура уже есть/предусмотрена")
            has_files = st.checkbox("Есть исходные файлы")
            new_project = st.checkbox("Проект новый (с нуля)")
        volume = st.number_input("Объём добычи, т/год", value=500000, step=50000)
        submitted = st.form_submit_button("➡️ Создать проект")
    if submitted and name:
        S.project = {"name": name, "underground": underground, "open_pit": open_pit,
                     "blasting": blasting, "infrastructure": infrastructure,
                     "has_files": has_files, "new_project": new_project, "volume": volume}
        S.wizard_done = True
        log_request(user_id, "wizard", json.dumps(S.project, ensure_ascii=False))
        st.success(f"Проект «{name}» создан! Данные подтянет «🏗️ Цифровой штаб».")
    elif S.wizard_done:
        p = S.project
        st.info(f"Текущий проект: «{p['name']}» — {p['volume']:,} т/год".replace(",", " "))
        if st.button("🔄 Начать новый проект"):
            S.wizard_done = False
            S.docs = []
            S.project = None
            st.rerun()

# ---------- ДОКУМЕНТЫ ----------
elif menu.startswith("📂"):
    st.header("Шаг 2: загрузка документов проекта")
    if not S.wizard_done:
        st.warning("Сначала создайте проект в разделе «🚀 Новый проект».")
    else:
        files = st.file_uploader("Выберите файлы", type=["pdf", "docx", "txt", "md", "csv", "xlsx", "dwg"],
                                 accept_multiple_files=True)
        if files:
            for f in files:
                text = extract_text(f)
                types = classify_doc(f.name, text)
                safe_name = f"{datetime.datetime.now():%Y%m%d%H%M%S}_{f.name}"
                try:
                    with open(os.path.join(DOCS_DIR, safe_name), "wb") as out:
                        out.write(f.getvalue())
                except Exception:
                    pass
                S.docs.append({"name": f.name, "text": text, "types": types})
                log_request(user_id, "upload", f.name, ",".join(types))
            st.success(f"Загружено файлов: {len(files)}. База знаний обновлена.")
        if S.docs:
            st.subheader(f"В базе знаний: {len(S.docs)} документов")
            for d in S.docs:
                with st.expander(f"📄 {d['name']} — {', '.join(d['types'])}"):
                    st.text(d["text"][:1500] or "[текст не извлечён]")

# ---------- ЧЕК-ЛИСТ ----------
elif menu.startswith("✅"):
    st.header("Шаг 3: чек-лист готовности проекта")
    if not S.wizard_done:
        st.warning("Сначала создайте проект в разделе «🚀 Новый проект».")
    else:
        present = set()
        for d in S.docs:
            present.update(d["types"])
        base, extra = readiness_check(present, S.project)
        ok = sum(1 for _, s, _ in base if s == "ok")
        st.progress(ok / len(base) if base else 0, f"Готовность: {ok} из {len(base)}")
        for dt, status, advice in base:
            if status == "ok":
                st.markdown(f"✅ **{dt}** — загружено")
            else:
                st.markdown(f"⚠️ **{dt}** — отсутствует")
                if advice:
                    st.markdown(f"&nbsp;&nbsp;&nbsp;→ {advice}")
        if extra:
            st.subheader("Отраслевые требования")
            for dt, status, advice in extra:
                st.markdown(f"⚠️ **{dt}** — требуется")
                st.markdown(f"&nbsp;&nbsp;&nbsp;→ {advice}")

# ---------- ИНЖЕНЕРНАЯ ЛОГИКА ----------
elif menu.startswith("🧠"):
    st.header("Инженерная логика: требования к проекту")
    if not S.wizard_done:
        st.warning("Сначала создайте проект.")
    elif st.button("Построить цепочку требований"):
        log_request(user_id, "logic", json.dumps(S.project, ensure_ascii=False))
        for i, r in enumerate(build_requirements(S.project), 1):
            st.markdown(f"{i}. {r}")

# ---------- АССИСТЕНТ ----------
elif menu.startswith("💬"):
    st.header("Ассистент по базе знаний проекта")
    if not S.docs:
        st.info("База знаний пуста. Загрузите документы.")
    else:
        q = st.text_input("Ваш вопрос")
        if st.button("Спросить") and q.strip():
            log_request(user_id, "ask", q)
            results = search_knowledge(q)
            if results:
                for score, fname, frag in results:
                    st.markdown(f"**📄 {fname}:**")
                    st.info(frag.strip()[:800])
            else:
                st.warning("Релевантная информация не найдена.")

# ---------- УСТАВНЫЕ ----------
elif menu.startswith("📋"):
    st.header("Уставные документы")
    if not S.wizard_done:
        st.warning("Сначала создайте проект.")
    else:
        doc = st.radio("Документ", ["Устав", "График", "Бюджет"])
        if st.button("Сгенерировать"):
            log_request(user_id, "charter", doc)
            fn = {"Устав": gen_charter, "График": gen_schedule, "Бюджет": gen_budget}[doc]
            st.code(fn(S.project), language="text")

# ---------- DWG ----------
elif menu.startswith("📐"):
    st.header("PRO: анализ DWG-чертежей")
    if S.plan != "pro":
        st.warning(f"🔒 Тариф PRO ({PRICE_PRO} ₽/мес).")
    else:
        f = st.file_uploader("Загрузите DWG-файл", type=["dwg"])
        if f and st.button("📐 Проанализировать чертёж"):
            log_request(user_id, "dwg", f.name)
            text = extract_text(f)
            st.markdown("**Рекомендации по оптимизации:**")
            st.markdown(analyze_dwg(text, f.name))

# ---------- ПОДПИСКА ----------
elif menu.startswith("💳"):
    st.header("Подписка")
    st.markdown(f"**Базовая** — {PRICE_BASIC} ₽/мес.")
    st.markdown(f"**PRO** — {PRICE_PRO} ₽/мес: + DWG-анализ.")
    promo = st.text_input("Промокод").strip().upper()
    if st.button("Активировать"):
        log_request(user_id, "billing", f"promo={promo}")
        if promo in PROMO_CODES:
            S.plan = PROMO_CODES[promo]
            st.success(f"Активирован тариф: {PROMO_CODES[promo].upper()}")
        else:
            st.error("Неверный промокод.")

# ---------- КАБИНЕТ ВЛАДЕЛЬЦА ----------
elif menu.startswith("🔐"):
    st.header("Кабинет владельца")
    if not S.show_owner:
        pwd = st.text_input("Пароль владельца", type="password")
        if st.button("Войти") and pwd == OWNER_PASSWORD:
            S.show_owner = True
            st.rerun()
    if S.show_owner:
        entries = read_log()
        st.subheader(f"Журнал: {len(entries)} записей")
        users = {}
        for e in entries:
            users.setdefault(e["user"], []).append(e)
        for u, es in users.items():
            with st.expander(f"{u} — {len(es)} действий"):
                st.markdown("**Психологический портрет:**")
                st.markdown(psych_profile(es))
                for e in es[-20:]:
                    st.text(f"[{e['ts']}] ({e['mode']}) {e['text'][:120]}")
        st.subheader("База знаний (файлы на сервере)")
        for fn in sorted(os.listdir(DOCS_DIR)):
            st.text(fn)

st.sidebar.markdown("---")
st.sidebar.caption("© 2026 [Владелец продукта]. Все права защищены.")
